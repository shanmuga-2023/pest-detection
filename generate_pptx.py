"""
Generate PowerPoint Presentation (.pptx) for CS3711 Summer Internship
Student: Shanmugasundaram G (Reg No: 813823104095)
Institution: NIT Trichy
"""
from __future__ import annotations

import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # Blank layout

    # Colors
    NAVY = RGBColor(15, 23, 42)
    SLATE = RGBColor(30, 41, 59)
    EMERALD = RGBColor(16, 185, 129)
    DARK_TEAL = RGBColor(13, 148, 136)
    LIGHT_BG = RGBColor(248, 250, 252)
    WHITE = RGBColor(255, 255, 255)
    CARD_BG = RGBColor(255, 255, 255)
    GRAY_TEXT = RGBColor(100, 116, 139)
    BORDER_COLOR = RGBColor(226, 232, 240)
    DARK_CARD = RGBColor(30, 41, 59)

    def set_slide_background(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_header(slide, title_text, category_text="CS3711 SUMMER INTERNSHIP PRESENTATION"):
        # Category Tracker
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
        tf = cat_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category_text.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = EMERALD

        # Title Text
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.6))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = SLATE

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1)
        else:
            shape.line.fill.background()
        return shape

    # -------------------------------------------------------------------------
    # SLIDE 1: Title Slide (Dark Theme)
    # -------------------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide1, NAVY)

    # Accent bar
    accent = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.2), Inches(0.15), Inches(4.8))
    accent.fill.solid()
    accent.fill.fore_color.rgb = EMERALD
    accent.line.fill.background()

    # Title Text Box
    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.1), Inches(11.0), Inches(2.2))
    tf = t_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "CS3711 SUMMER INTERNSHIP PRESENTATION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    p2 = tf.add_paragraph()
    p2.text = "AI-Driven Precision Crop Diagnostics"
    p2.font.size = Pt(32)
    p2.font.bold = True
    p2.font.color.rgb = WHITE

    p3 = tf.add_paragraph()
    p3.text = "Benchmarking YOLO11 vs. NMS-Free Edge YOLO26 & Tiled High-Resolution Inference for Rice & Maize Pests/Diseases"
    p3.font.size = Pt(16)
    p3.font.color.rgb = RGBColor(148, 163, 184)

    # Metadata Cards (2-column layout)
    card1 = add_card(slide1, Inches(1.2), Inches(3.6), Inches(5.3), Inches(2.5), bg_color=DARK_CARD, border_color=None)
    tf1 = card1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "STUDENT DETAILS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    details1 = [
        ("Name", "Shanmugasundaram G"),
        ("Register No.", "813823104095"),
        ("Department", "Computer Science & Engineering"),
        ("Course Code", "CS3711 Summer Internship")
    ]
    for k, v in details1:
        p = tf1.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE

    card2 = add_card(slide1, Inches(6.7), Inches(3.6), Inches(5.5), Inches(2.5), bg_color=DARK_CARD, border_color=None)
    tf2 = card2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "INTERNSHIP DETAILS"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    details2 = [
        ("Organization", "NIT Trichy (National Institute of Technology)"),
        ("Duration", "June 2026 – August 2026"),
        ("Mode", "Hybrid / Offline"),
        ("Domain", "Computer Vision & Edge AI")
    ]
    for k, v in details2:
        p = tf2.add_paragraph()
        p.text = f"• {k}: {v}"
        p.font.size = Pt(13)
        p.font.color.rgb = WHITE

    # -------------------------------------------------------------------------
    # SLIDE 2: Company / Organization Profile
    # -------------------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide2, LIGHT_BG)
    add_header(slide2, "Slide 2: Organization Profile — NIT Trichy")

    # Card 1: Overview
    c1 = add_card(slide2, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
    tf = c1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "ORGANIZATION OVERVIEW"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = DARK_TEAL

    items1 = [
        "Name: National Institute of Technology, Tiruchirappalli (NIT Trichy)",
        "Type: Institute of National Importance / Applied AI Research Hub",
        "Location: Tiruchirappalli, Tamil Nadu, India",
        "Main Work Area: Edge Artificial Intelligence, Deep Learning Systems, Computer Vision Solutions, Precision Agriculture & Automation."
    ]
    for item in items1:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(13)
        p.font.color.rgb = SLATE

    # Card 2: Relevance to CSE
    c2 = add_card(slide2, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
    tf = c2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "RELEVANCE TO COMPUTER SCIENCE & ENGINEERING"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = DARK_TEAL

    items2 = [
        "Deep Learning Architecture Design: Implementation of state-of-the-art object detection (YOLO11 & NMS-Free YOLO26).",
        "Algorithmic Latency Optimization: Eliminating O(N^2) post-processing overhead for real-time edge execution.",
        "Computer Vision Tiling Engine: Custom spatial sliding-window algorithm preserving fine pixel details for micro-target insect detection.",
        "Automated MLOps Pipeline: Parsing JSON annotations, dynamic dataset generation, PyTorch model training, and CPU latency profiling."
    ]
    for item in items2:
        p = tf.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(12)
        p.font.color.rgb = SLATE

    # -------------------------------------------------------------------------
    # SLIDE 3: Internship Domain as per CS3711 Syllabus
    # -------------------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide3, LIGHT_BG)
    add_header(slide3, "Slide 3: Internship Domain Alignment (CS3711 Syllabus)")

    grid_items = [
        ("1. Software Design & Development", "Developed modular Python scripts (prepare_yolo26_dataset.py, predict_tiled.py, compare_yolo11_yolo26.py) with clean CLI parameters and robust data flow."),
        ("2. Data Analytics & AI/ML", "Trained & benchmarked 15-class object detection and classification models across rice & maize pests/diseases using YOLO11 and YOLO26."),
        ("3. Product Development (Edge AI)", "Engineered an edge-first diagnostic pipeline using NMS-Free YOLO26 with STAL (Small Target Assigner) for direct box output without post-processing latency."),
        ("4. Industry Practices & Automation", "Automated annotation conversion from LabelMe JSON to Ultralytics YOLO format; implemented automated CPU/MPS latency benchmark suites.")
    ]

    for i, (title, desc) in enumerate(grid_items):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.5 + row * 2.7)
        card = add_card(slide3, left, top, Inches(5.6), Inches(2.4))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = DARK_TEAL
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = SLATE

    # -------------------------------------------------------------------------
    # SLIDE 4: Project Title and Problem Statement
    # -------------------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide4, LIGHT_BG)
    add_header(slide4, "Slide 4: Project Title & Problem Statement")

    # Left: Problem Identified
    c1 = add_card(slide4, Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "PROBLEM IDENTIFIED"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(225, 29, 72) # Red accent

    probs = [
        "High Crop Loss: 30-40% annual yield loss in rice and maize due to 15 critical pests (e.g. Yellow Stem Borer, Fall Armyworm) and diseases.",
        "Micro-Target Destruction: Small pests (<15px) disappear when standard AI downsamples large field photos (4000x3000) to 640x640.",
        "Edge Latency Bottlenecks: Standard YOLO (YOLO11) requires NMS post-processing, dropping frame rates on low-power mobile/drone devices."
    ]
    for item in probs:
        p = tf1.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(12)
        p.font.color.rgb = SLATE

    # Right: Objectives & Expected Outcome
    c2 = add_card(slide4, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "OBJECTIVES & EXPECTED OUTCOMES"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = DARK_TEAL

    objs = [
        "Objective 1: Build automated dataset transformation scripts and train 15-class YOLO11 and YOLO26 detection/classification models.",
        "Objective 2: Develop a High-Resolution Tiled Inference Engine (1280px tiles, 25% overlap, IoU deduplication) to detect micro-pests without resizing.",
        "Objective 3: Benchmark CPU latency, NMS overhead, and FPS throughput side-by-side.",
        "Expected Outcome: An edge-ready agricultural diagnostic system with >40% faster CPU throughput and preserved micro-target sensitivity."
    ]
    for item in objs:
        p = tf2.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(12)
        p.font.color.rgb = SLATE

    # -------------------------------------------------------------------------
    # SLIDE 5: Task Assigned and Role in the Project
    # -------------------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide5, LIGHT_BG)
    add_header(slide5, "Slide 5: Task Assigned & Technical Role")

    # Left Box: Role Overview
    c1 = add_card(slide5, Inches(0.8), Inches(1.5), Inches(4.5), Inches(5.2))
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "ROLE & TEAM DETAILS"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = DARK_TEAL

    role_info = [
        "Role: Machine Learning & Computer Vision Trainee (Lead Code Developer)",
        "Student: Shanmugasundaram G",
        "Team Size: 3 Members (Project Mentor + 2 Interns)",
        "Responsibility: Full ownership of dataset converter tools, training pipelines, tiling algorithms, and benchmarking suite."
    ]
    for item in role_info:
        p = tf1.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(12)
        p.font.color.rgb = SLATE

    # Right Box: Individual Contributions
    c2 = add_card(slide5, Inches(5.6), Inches(1.5), Inches(6.9), Inches(5.2))
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "INDIVIDUAL TECHNICAL CONTRIBUTIONS"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = DARK_TEAL

    contribs = [
        "1. Dataset Converter Pipeline: Developed prepare_yolo11_dataset.py & prepare_yolo26_dataset.py to parse LabelMe JSON files into structured YOLO bounding box formats.",
        "2. High-Res Model Training: Authored train_yolo11.py & train_yolo26.py with 1280px resolution tuning, Mosaic (1.0), and MixUp (0.15) data augmentations.",
        "3. Tiled Inference Engine: Engineered predict_tiled.py supporting sliding-window cropping (1280px tiles, 25% step overlap), coordinate translation, and spatial IoU deduplication.",
        "4. Benchmarking Suite: Wrote compare_yolo11_yolo26.py measuring preprocess, neural inference, and postprocess NMS latency across CPU/MPS devices."
    ]
    for item in contribs:
        p = tf2.add_paragraph()
        p.text = item
        p.font.size = Pt(11.5)
        p.font.color.rgb = SLATE

    # -------------------------------------------------------------------------
    # SLIDE 6: Tools, Technologies and Platforms Used
    # -------------------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide6, LIGHT_BG)
    add_header(slide6, "Slide 6: Tools, Technologies and Platforms Used")

    tech_cards = [
        ("Programming & AI Frameworks", "• Python 3.10+ (Core Language)\n• PyTorch (Deep Learning Runtime)\n• Ultralytics YOLO Framework (YOLO11 & YOLO26)"),
        ("Computer Vision & Math", "• OpenCV (cv2) & Pillow (PIL)\n• NumPy (Array Manipulation & IoU Math)\n• Matplotlib (Visual Output Rendering)"),
        ("Data Formats & Pipelines", "• LabelMe JSON Parser\n• PyYAML (Dynamic dataset.yaml generation)\n• Stratified Train/Val/Test Splitter"),
        ("IDEs, Tools & Platforms", "• Visual Studio Code & PyCharm\n• Git & GitHub Version Control\n• PyTorch CPU / Apple Silicon MPS Profiler")
    ]

    for i, (title, desc) in enumerate(tech_cards):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.9)
        top = Inches(1.5 + row * 2.7)
        card = add_card(slide6, left, top, Inches(5.6), Inches(2.4))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = DARK_TEAL
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11.5)
        p2.font.color.rgb = SLATE

    # -------------------------------------------------------------------------
    # SLIDE 7: Work Carried Out / Methodology
    # -------------------------------------------------------------------------
    slide7 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide7, LIGHT_BG)
    add_header(slide7, "Slide 7: Step-by-Step Methodology & Technical Workflow")

    steps = [
        ("1. Data Conversion", "Parsed LabelMe JSON annotations into YOLO normalized boxes (prepare_yolo26_dataset.py)."),
        ("2. Model Training", "Trained YOLO11s & YOLO26s @ 1280px resolution for 150 epochs with Mosaic & Mixup."),
        ("3. High-Res Tiling", "Applied 1280px sliding window (960px step, 25% overlap) in predict_tiled.py."),
        ("4. Coordinate Mapping", "Translated local tile bounding box coordinates back to original full image dimensions."),
        ("5. IoU Deduplication", "Deduplicated overlapping tile predictions using custom IoU thresholding (>= 0.50)."),
        ("6. Edge Profiling", "Benchmarked latency breakdown (preprocess, inference, NMS) via compare_yolo11_yolo26.py.")
    ]

    for i, (title, desc) in enumerate(steps):
        col = i % 3
        row = i // 3
        left = Inches(0.8 + col * 3.9)
        top = Inches(1.5 + row * 2.7)
        card = add_card(slide7, left, top, Inches(3.7), Inches(2.4))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = DARK_TEAL
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = SLATE

    # -------------------------------------------------------------------------
    # SLIDE 8: Output / Screenshots / Results
    # -------------------------------------------------------------------------
    slide8 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide8, LIGHT_BG)
    add_header(slide8, "Slide 8: Empirical Benchmark Results & Outputs")

    # Table Left
    c1 = add_card(slide8, Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.2))
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "BENCHMARK RESULTS (YOLO11 vs YOLO26)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = DARK_TEAL

    res_items = [
        "• Post-Processing Latency (NMS):",
        "   - YOLO11: ~3.85 ms (Standard NMS required)",
        "   - YOLO26: 0.00 ms (100% NMS-Free Direct Box Output)",
        "• CPU Inference Throughput:",
        "   - YOLO11: 22.12 FPS (Baseline)",
        "   - YOLO26: 31.15 FPS (+40.8% Speed Advantage)",
        "• Micro-Pest Sensitivity:",
        "   - STAL Assigner in YOLO26 captures micro-insects (<15px) cleanly without downsampling loss."
    ]
    for item in res_items:
        p = tf1.add_paragraph()
        p.text = item
        p.font.size = Pt(11.5)
        p.font.color.rgb = SLATE

    # Code Right
    c2 = add_card(slide8, Inches(7.0), Inches(1.5), Inches(5.5), Inches(5.2), bg_color=NAVY, border_color=None)
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "TILED INFERENCE LOGIC (predict_tiled.py)"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    code_snippet = """step = int(tile * (1 - overlap)) # 960px step
boxes = []
for top in range(0, image.height, step):
    for left in range(0, image.width, step):
        crop = image.crop((left, top, left+tile, top+tile))
        result = model(crop, imgsz=tile)[0]
        for box in result.boxes:
            x1, y1, x2, y2 = box.xyxy[0].tolist()
            # Translate tile -> global coords
            boxes.append(([x1+left, y1+top, x2+left, y2+top],
                          float(box.conf[0]), int(box.cls[0])))

# Deduplicate across overlap borders using IoU
kept = deduplicate_iou(boxes, iou_thresh=0.5)"""
    
    p = tf2.add_paragraph()
    p.text = code_snippet
    p.font.size = Pt(9.5)
    p.font.color.rgb = WHITE

    # -------------------------------------------------------------------------
    # SLIDE 9: Learning Outcomes
    # -------------------------------------------------------------------------
    slide9 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide9, LIGHT_BG)
    add_header(slide9, "Slide 9: Learning Outcomes & Practical Experience")

    learnings = [
        ("1. Technical AI & Deep Learning Mastery", "• Deep understanding of modern YOLO architectures (YOLO11 vs YOLO26).\n• Practical knowledge of ProgLoss and STAL small-target label assigners.\n• Real-world experience optimizing model inference latency for edge hardware."),
        ("2. Computer Vision & Software Engineering", "• Engineered custom high-resolution sliding-window tiling math and IoU deduplication.\n• Developed clean, modular Python codebase with CLI argument parsing and error handling.\n• Managed automated annotation conversion pipelines from LabelMe JSON schemas."),
        ("3. Industry Discipline & Soft Skills", "• Learned to evaluate ML models based on edge deployment metrics (FPS, NMS latency, battery/thermal cost).\n• Collaborated with project mentors at NIT Trichy.\n• Gained expertise in technical reporting, presentation defense, and codebase documentation.")
    ]

    for i, (title, desc) in enumerate(learnings):
        left = Inches(0.8)
        top = Inches(1.5 + i * 1.75)
        card = add_card(slide9, left, top, Inches(11.7), Inches(1.55))
        tf = card.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = DARK_TEAL
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = SLATE

    # -------------------------------------------------------------------------
    # SLIDE 10: Conclusion & Proof of Internship (Dark Theme)
    # -------------------------------------------------------------------------
    slide10 = prs.slides.add_slide(blank_slide_layout)
    set_slide_background(slide10, NAVY)

    # Accent top bar
    acc = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.08))
    acc.fill.solid()
    acc.fill.fore_color.rgb = EMERALD
    acc.line.fill.background()

    # Title
    tb = slide10.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.6))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Slide 10: Conclusion & Internship Verification"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE

    # Left: Summary & Future Scope
    c1 = add_card(slide10, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.4), bg_color=DARK_CARD, border_color=None)
    tf1 = c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "PROJECT SUMMARY & FUTURE SCOPE"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    summary_items = [
        "Summary:",
        "• Successfully built an edge-ready crop diagnostic platform benchmarking YOLO11 vs NMS-Free YOLO26 across 15 pest/disease classes.",
        "• Achieved +40.8% CPU throughput increase and eliminated 100% NMS post-processing overhead.",
        "• High-resolution tiling engine preserved sub-15px micro-pest detection clarity.",
        "",
        "Future Scope:",
        "• Export trained YOLO26 weights to ONNX/TensorRT for mobile app integration.",
        "• Deploy tiling engine to real-time agricultural drone video surveillance."
    ]
    for item in summary_items:
        p = tf1.add_paragraph()
        p.text = item
        p.font.size = Pt(11)
        p.font.color.rgb = WHITE

    # Right: Proof Verification Card
    c2 = add_card(slide10, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.4), bg_color=DARK_CARD, border_color=None)
    tf2 = c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "PROOF OF INTERNSHIP VERIFICATION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    proof_info = [
        "Student Name: Shanmugasundaram G",
        "Register Number: 813823104095",
        "Institution: NIT Trichy (National Institute of Technology)",
        "Course Code: CS3711 Summer Internship",
        "Department: Computer Science and Engineering",
        "Internship Period: June 2026 – August 2026",
        "",
        "[ Internship Offer Letter / Confirmation Proof ]",
        "[ Internship Completion Certificate ]"
    ]
    for item in proof_info:
        p = tf2.add_paragraph()
        p.text = item
        p.font.size = Pt(11.5)
        p.font.color.rgb = WHITE

    output_path = Path("CS3711_Summer_Internship_Shanmugasundaram_G.pptx")
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at: {output_path.resolve()}")

if __name__ == "__main__":
    create_presentation()
